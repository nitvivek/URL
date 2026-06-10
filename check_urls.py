#!/usr/bin/env python3
"""
Check all distinct URLs in column E of the Excel file for HTTP accessibility.
Produces url_check_results.json with categorized results.
"""

import json
import time
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed

import openpyxl
import requests
from requests.exceptions import (
    ConnectionError,
    SSLError,
    Timeout,
    TooManyRedirects,
)
from urllib3.exceptions import InsecureRequestWarning

# Suppress SSL warnings for verification-disabled requests
requests.packages.urllib3.disable_warnings(InsecureRequestWarning)

EXCEL_FILE = "Master list - Ministries and Departments - PILOT BATCH (rows 32-42) (1).xlsx"
SHEET_NAME = "Main"
URL_COLUMN = 5  # Column E
TIMEOUT = 15
MAX_REDIRECTS = 5
MAX_WORKERS = 10
USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/125.0.0.0 Safari/537.36"
)


def get_urls_with_rows(wb_path, sheet_name):
    """Read all URLs from column E and map each URL to its row numbers."""
    wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)
    ws = wb[sheet_name]

    url_to_rows = defaultdict(list)
    for row_idx in range(2, ws.max_row + 1):
        cell_value = ws.cell(row=row_idx, column=URL_COLUMN).value
        if cell_value and str(cell_value).strip():
            url = str(cell_value).strip()
            url_to_rows[url].append(row_idx)

    wb.close()
    return url_to_rows


def classify_error(exc):
    """Classify an exception into an error type string."""
    if isinstance(exc, Timeout):
        return "timeout"
    elif isinstance(exc, SSLError):
        return "ssl"
    elif isinstance(exc, ConnectionError):
        err_str = str(exc).lower()
        if "name" in err_str and "resolution" in err_str:
            return "dns"
        if "nodename nor servname" in err_str:
            return "dns"
        if "getaddrinfo" in err_str:
            return "dns"
        if "nameresolutionerror" in err_str:
            return "dns"
        return "connection"
    elif isinstance(exc, TooManyRedirects):
        return "too_many_redirects"
    else:
        return "unknown"


def check_url(url):
    """
    Check a single URL. Try HEAD first, fall back to GET.
    Returns a dict with the result.
    """
    headers = {"User-Agent": USER_AGENT}
    session = requests.Session()
    session.max_redirects = MAX_REDIRECTS

    # Try HEAD first
    for method in [session.head, session.get]:
        try:
            response = method(
                url,
                headers=headers,
                timeout=TIMEOUT,
                allow_redirects=True,
                verify=True,
            )
            status_code = response.status_code

            if 200 <= status_code < 400:
                return {
                    "url": url,
                    "status": "working",
                    "status_code": status_code,
                    "final_url": response.url,
                }
            elif method == session.head and status_code >= 400:
                # HEAD might be blocked, try GET
                continue
            else:
                # GET also returned error
                error_type = f"{status_code // 100}xx"
                return {
                    "url": url,
                    "status": "broken",
                    "status_code": status_code,
                    "error_type": error_type,
                }
        except SSLError:
            # Retry with SSL verification disabled
            try:
                response = method(
                    url,
                    headers=headers,
                    timeout=TIMEOUT,
                    allow_redirects=True,
                    verify=False,
                )
                status_code = response.status_code
                if 200 <= status_code < 400:
                    return {
                        "url": url,
                        "status": "working",
                        "status_code": status_code,
                        "final_url": response.url,
                        "note": "SSL certificate invalid but site accessible",
                    }
                elif method == session.head and status_code >= 400:
                    continue
                else:
                    error_type = f"{status_code // 100}xx"
                    return {
                        "url": url,
                        "status": "broken",
                        "status_code": status_code,
                        "error_type": error_type,
                    }
            except Exception as e2:
                if method == session.head:
                    continue
                return {
                    "url": url,
                    "status": "broken",
                    "status_code": None,
                    "error_type": "ssl",
                }
        except (Timeout, ConnectionError, TooManyRedirects) as e:
            if method == session.head:
                # Try GET as fallback
                continue
            return {
                "url": url,
                "status": "broken",
                "status_code": None,
                "error_type": classify_error(e),
            }
        except Exception as e:
            if method == session.head:
                continue
            return {
                "url": url,
                "status": "broken",
                "status_code": None,
                "error_type": "unknown",
            }

    # Should not reach here, but just in case
    return {
        "url": url,
        "status": "broken",
        "status_code": None,
        "error_type": "unknown",
    }


def main():
    print(f"Loading URLs from: {EXCEL_FILE}")
    url_to_rows = get_urls_with_rows(EXCEL_FILE, SHEET_NAME)
    total_urls = len(url_to_rows)
    print(f"Found {total_urls} distinct URLs to check")
    print(f"Using {MAX_WORKERS} concurrent workers with {TIMEOUT}s timeout\n")

    working = []
    broken = []
    checked = 0
    start_time = time.time()

    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
        future_to_url = {
            executor.submit(check_url, url): url for url in url_to_rows.keys()
        }

        for future in as_completed(future_to_url):
            result = future.result()
            checked += 1
            url = result["url"]

            if result["status"] == "working":
                working.append(url)
                status_str = f"OK ({result['status_code']})"
            else:
                broken_entry = {
                    "url": url,
                    "error_type": result["error_type"],
                    "status_code": result.get("status_code"),
                    "rows_affected": url_to_rows[url],
                }
                broken.append(broken_entry)
                status_str = f"BROKEN ({result['error_type']}, code={result.get('status_code')})"

            if checked % 10 == 0 or checked == total_urls:
                elapsed = time.time() - start_time
                print(f"  [{checked}/{total_urls}] ({elapsed:.1f}s) Last: {url[:60]}... -> {status_str}")

    elapsed_total = time.time() - start_time

    # Sort broken by number of affected rows (most impactful first)
    broken.sort(key=lambda x: len(x["rows_affected"]), reverse=True)

    # Build result
    results = {
        "working": sorted(working),
        "broken": broken,
        "summary": {
            "total": total_urls,
            "working": len(working),
            "broken": len(broken),
            "check_duration_seconds": round(elapsed_total, 1),
        },
    }

    # Save results
    output_file = "url_check_results.json"
    with open(output_file, "w", encoding="utf-8") as f:
        json.dump(results, f, indent=2, ensure_ascii=False)

    # Print summary
    print(f"\n{'='*60}")
    print(f"URL CHECK COMPLETE")
    print(f"{'='*60}")
    print(f"Total URLs checked: {total_urls}")
    print(f"Working (2xx/3xx):  {len(working)}")
    print(f"Broken:             {len(broken)}")
    print(f"Time taken:         {elapsed_total:.1f} seconds")
    print(f"\nResults saved to: {output_file}")

    if broken:
        print(f"\n{'='*60}")
        print(f"BROKEN URLs (sorted by impact):")
        print(f"{'='*60}")
        for entry in broken[:20]:  # Show top 20
            rows_str = ", ".join(str(r) for r in entry["rows_affected"][:5])
            if len(entry["rows_affected"]) > 5:
                rows_str += f"... (+{len(entry['rows_affected'])-5} more)"
            print(f"  [{entry['error_type']}] {entry['url']}")
            print(f"    Status: {entry['status_code']}, Rows: {rows_str}")
        if len(broken) > 20:
            print(f"\n  ... and {len(broken)-20} more broken URLs (see {output_file})")


if __name__ == "__main__":
    main()
